# # from fastapi import APIRouter, HTTPException
# # from pydantic import BaseModel
# # from solengineer.api import state
# # from solengineer.perf.perf_tester import PerfTester

# # router = APIRouter(prefix="/api")


# # class PerfRequest(BaseModel):
# #     topic: str
# #     message_count: int = 1000
# #     message_size: int = 256


# # @router.post("/perf")
# # def run_perf(req: PerfRequest):
# #     service = state.get_service()
# #     if not service:
# #         raise HTTPException(status_code=400, detail="Not connected to broker")
# #     tester = PerfTester(service)
# #     results = tester.run(req.topic, req.message_count, req.message_size)
# #     return results


# #updated with some extra featueres


# # from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect
# # from pydantic import BaseModel
# # from solengineer.api import state
# # from solengineer.perf.perf_tester import PerfTester
# # import asyncio
# # import threading
# # import json
# # from datetime import datetime

# # # HTTP routes use /api prefix, WebSocket does NOT
# # router = APIRouter()

# # _current_tester: PerfTester | None = None


# # class PerfRequest(BaseModel):
# #     target: str
# #     target_type: str = "topic"
# #     message_count: int = 1000
# #     message_size: int = 256
# #     message_template: str = ""


# # @router.post("/api/perf")
# # def run_perf(req: PerfRequest):
# #     service = state.get_service()
# #     if not service:
# #         raise HTTPException(status_code=400, detail="Not connected to broker")
# #     tester = PerfTester(service)
# #     results = tester.run(
# #         target=req.target,
# #         target_type=req.target_type,
# #         message_count=req.message_count,
# #         message_size=req.message_size,
# #         message_template=req.message_template
# #     )
# #     return results


# # @router.post("/api/perf/stop")
# # def stop_perf():
# #     global _current_tester
# #     if _current_tester:
# #         _current_tester.stop()
# #     return {"status": "stopped"}


# # @router.websocket("/ws/perf")
# # async def perf_ws(websocket: WebSocket):
# #     global _current_tester
# #     await websocket.accept()

# #     service = state.get_service()
# #     if not service:
# #         await websocket.send_json({"error": "Not connected to broker"})
# #         await websocket.close()
# #         return

# #     try:
# #         config = await asyncio.wait_for(websocket.receive_text(), timeout=10)
# #         req = json.loads(config)
# #     except Exception:
# #         await websocket.close()
# #         return

# #     tester = PerfTester(service)
# #     _current_tester = tester
# #     results_holder = {}
# #     loop = asyncio.get_event_loop()

# #     def on_progress(received, total, elapsed, tps):
# #         asyncio.run_coroutine_threadsafe(
# #             websocket.send_json({
# #                 "type": "progress",
# #                 "received": received,
# #                 "total": total,
# #                 "elapsed": elapsed,
# #                 "tps": tps,
# #                 "percent": round(received / total * 100, 1)
# #             }),
# #             loop
# #         )

# #     def run_test():
# #         try:
# #             results_holder["data"] = tester.run(
# #                 target=req.get("target", "perf/test"),
# #                 target_type=req.get("target_type", "topic"),
# #                 message_count=req.get("message_count", 1000),
# #                 message_size=req.get("message_size", 256),
# #                 message_template=req.get("message_template", ""),
# #                 on_progress=on_progress
# #             )
# #         except Exception as e:
# #             results_holder["error"] = str(e)

# #     thread = threading.Thread(target=run_test)
# #     thread.start()

# #     try:
# #         while thread.is_alive():
# #             await asyncio.sleep(0.3)
# #         thread.join()

# #         if "error" in results_holder:
# #             await websocket.send_json({
# #                 "type": "error",
# #                 "error": results_holder["error"]
# #             })
# #         elif "data" in results_holder:
# #             await websocket.send_json({
# #                 "type": "complete",
# #                 "results": results_holder["data"],
# #                 "timestamp": datetime.now().isoformat()
# #             })
# #     except WebSocketDisconnect:
# #         tester.stop()
# #     finally:
# #         try:
# #             await websocket.close()
# #         except Exception:
# #             pass


# # 21-06-2026

# from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect
# from pydantic import BaseModel
# from solengineer.api import state
# from solengineer.perf.perf_tester import PerfTester
# import asyncio
# import threading
# import json
# from datetime import datetime

# router = APIRouter()

# _current_tester: PerfTester | None = None


# class PerfRequest(BaseModel):
#     target: str
#     target_type: str = "topic"
#     message_count: int = 1000
#     message_size: int = 256
#     message_template: str = ""
#     delivery_mode: str = "direct"
#     rate_limit: int = 0
#     window_size: int = 50


# @router.post("/api/perf")
# def run_perf(req: PerfRequest):
#     service = state.get_service()
#     if not service:
#         raise HTTPException(status_code=400, detail="Not connected to broker")
#     tester = PerfTester(service)
#     results = tester.run(
#         target=req.target,
#         target_type=req.target_type,
#         message_count=req.message_count,
#         message_size=req.message_size,
#         message_template=req.message_template,
#         delivery_mode=req.delivery_mode,
#         rate_limit=req.rate_limit,
#         window_size=req.window_size
#     )
#     return results


# @router.post("/api/perf/stop")
# def stop_perf():
#     global _current_tester
#     if _current_tester:
#         _current_tester.stop()
#     return {"status": "stopped"}


# @router.websocket("/ws/perf")
# async def perf_ws(websocket: WebSocket):
#     global _current_tester
#     await websocket.accept()

#     service = state.get_service()
#     if not service:
#         await websocket.send_json({"error": "Not connected to broker"})
#         await websocket.close()
#         return

#     try:
#         config = await asyncio.wait_for(websocket.receive_text(), timeout=10)
#         req = json.loads(config)
#     except Exception:
#         await websocket.close()
#         return

#     tester = PerfTester(service)
#     _current_tester = tester
#     results_holder = {}
#     loop = asyncio.get_event_loop()

#     def on_progress(received, total, elapsed, tps):
#         asyncio.run_coroutine_threadsafe(
#             websocket.send_json({
#                 "type": "progress",
#                 "received": received,
#                 "total": total,
#                 "elapsed": elapsed,
#                 "tps": tps,
#                 "percent": round(received / total * 100, 1)
#             }),
#             loop
#         )

#     def run_test():
#         try:
#             results_holder["data"] = tester.run(
#                 target=req.get("target", "perf/test"),
#                 target_type=req.get("target_type", "topic"),
#                 message_count=req.get("message_count", 1000),
#                 message_size=req.get("message_size", 256),
#                 message_template=req.get("message_template", ""),
#                 delivery_mode=req.get("delivery_mode", "direct"),
#                 rate_limit=req.get("rate_limit", 0),
#                 window_size=req.get("window_size", 50),
#                 on_progress=on_progress
#             )
#         except Exception as e:
#             results_holder["error"] = str(e)

#     thread = threading.Thread(target=run_test)
#     thread.start()

#     try:
#         while thread.is_alive():
#             await asyncio.sleep(0.3)
#         thread.join()

#         if "error" in results_holder:
#             await websocket.send_json({"type": "error", "error": results_holder["error"]})
#         elif "data" in results_holder:
#             await websocket.send_json({
#                 "type": "complete",
#                 "results": results_holder["data"],
#                 "timestamp": datetime.now().isoformat()
#             })
#     except WebSocketDisconnect:
#         tester.stop()
#     finally:
#         try:
#             await websocket.close()
#         except Exception:
#             pass

# # from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect
# # from fastapi.responses import StreamingResponse
# # from pydantic import BaseModel
# # from solengineer.api import state
# # from solengineer.perf.perf_tester import PerfTester
# # import asyncio
# # import threading
# # import json
# # import io
# # import openpyxl
# # from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
# # from openpyxl.chart import LineChart, Reference
# # from datetime import datetime

# # router = APIRouter()

# # _current_tester: PerfTester | None = None


# # class PerfRequest(BaseModel):
# #     target: str
# #     target_type: str = "topic"
# #     message_count: int = 1000
# #     message_size: int = 256
# #     message_template: str = ""
# #     delivery_mode: str = "direct"
# #     rate_limit: int = 0
# #     window_size: int = 50
# #     queue_type: str = "exclusive" # new 
# #     warmup_count: int = 0   # new 
# #     consumer_count: int = 1 #new 


# # @router.post("/api/perf")
# # def run_perf(req: PerfRequest):
# #     service = state.get_service()
# #     if not service:
# #         raise HTTPException(status_code=400, detail="Not connected to broker")
# #     tester = PerfTester(service)
# #     results = tester.run(
# #         target=req.target,
# #         target_type=req.target_type,
# #         message_count=req.message_count,
# #         message_size=req.message_size,
# #         message_template=req.message_template,
# #         delivery_mode=req.delivery_mode,
# #         rate_limit=req.rate_limit,
# #         window_size=req.window_size,
# #         queue_type=req.queue_type, # new 
# #         warmup_count=req.warmup_count # new 
        
# #     )
# #     return results


# # @router.post("/api/perf/stop")
# # def stop_perf():
# #     global _current_tester
# #     if _current_tester:
# #         _current_tester.stop()
# #     return {"status": "stopped"}


# # # ─── Excel Export ────────────────────────────────────────────────────────────

# # @router.post("/api/perf/export/excel")
# # def export_perf_excel(results: dict):

# #     wb = openpyxl.Workbook()

# #     # ── Styles ──────────────────────────────────────────────────
# #     header_font   = Font(bold=True, color="FFFFFF", size=11)
# #     header_fill   = PatternFill("solid", fgColor="1A73E8")
# #     label_font    = Font(bold=True, color="5A6675", size=10)
# #     value_font    = Font(color="1A2733", size=10)
# #     title_font    = Font(bold=True, color="1A2733", size=13)
# #     center        = Alignment(horizontal="center", vertical="center")
# #     left          = Alignment(horizontal="left", vertical="center")
# #     thin          = Side(style="thin", color="D4DAE0")
# #     border        = Border(left=thin, right=thin, top=thin, bottom=thin)

# #     # ── Sheet 1: Summary ────────────────────────────────────────
# #     ws = wb.active
# #     ws.title = "Summary"
# #     ws.sheet_view.showGridLines = False
# #     ws.column_dimensions["A"].width = 32
# #     ws.column_dimensions["B"].width = 32

# #     # Title row
# #     ws.merge_cells("A1:B1")
# #     ws["A1"] = "SolEngineer — Performance Test Results"
# #     ws["A1"].font = title_font
# #     ws["A1"].alignment = center
# #     ws["A1"].fill = PatternFill("solid", fgColor="F5F7FA")
# #     ws.row_dimensions[1].height = 32

# #     # Timestamp row
# #     ws.merge_cells("A2:B2")
# #     ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
# #     ws["A2"].font = Font(color="8A96A3", size=9, italic=True)
# #     ws["A2"].alignment = center
# #     ws.row_dimensions[2].height = 18

# #     def write_section(start_row, title, rows):
# #         # Section header
# #         ws.merge_cells(f"A{start_row}:B{start_row}")
# #         ws[f"A{start_row}"] = title
# #         ws[f"A{start_row}"].font = header_font
# #         ws[f"A{start_row}"].fill = header_fill
# #         ws[f"A{start_row}"].alignment = center
# #         ws[f"B{start_row}"].fill = header_fill
# #         ws.row_dimensions[start_row].height = 22
# #         r = start_row + 1
# #         for label, value in rows:
# #             ws.cell(r, 1, label).font = label_font
# #             ws.cell(r, 1).alignment = left
# #             ws.cell(r, 1).border = border
# #             ws.cell(r, 2, value).font = value_font
# #             ws.cell(r, 2).alignment = left
# #             ws.cell(r, 2).border = border
# #             ws.row_dimensions[r].height = 20
# #             r += 1
# #         return r

# #     # Configuration section
# #     next_row = write_section(4, "Configuration", [
# #         ("Target",              results.get("target", "")),
# #         ("Target Type",         results.get("target_type", "").capitalize()),
# #         ("Delivery Mode",       results.get("delivery_mode", "").capitalize()),
# #         ("Message Count",       results.get("message_count", "")),
# #         ("Message Size (bytes)", results.get("message_size_bytes", "")),
# #         ("Rate Limit (msg/s)",  results.get("rate_limit", 0) or "Unlimited"),
# #         ("Window Size",         results.get("window_size") or "N/A (Direct mode)"),
# #     ])

# #     # Results section
# #     next_row += 1
# #     write_section(next_row, "Results", [
# #         ("Throughput (msg/s)",  results.get("throughput_msg_per_sec", "")),
# #         ("Avg Latency (ms)",    results.get("avg_latency_ms", "")),
# #         ("P95 Latency (ms)",    results.get("p95_latency_ms", "")),
# #         ("P99 Latency (ms)",    results.get("p99_latency_ms", "")),
# #         ("Min Latency (ms)",    results.get("min_latency_ms", "")),
# #         ("Max Latency (ms)",    results.get("max_latency_ms", "")),
# #         ("Messages Sent",       results.get("sent", "")),
# #         ("Messages Received",   results.get("received", "")),
# #         ("Errors",              results.get("errors", "")),
# #         ("Total Time (sec)",    results.get("total_time_sec", "")),
# #     ])

# #     # ── Sheet 2: Throughput Over Time ────────────────────────────
# #     samples = results.get("throughput_samples", [])
# #     if samples:
# #         ws2 = wb.create_sheet("Throughput Over Time")
# #         ws2.sheet_view.showGridLines = False
# #         ws2.column_dimensions["A"].width = 15
# #         ws2.column_dimensions["B"].width = 25

# #         # Headers
# #         for col, text in [(1, "Second"), (2, "Throughput (msg/s)")]:
# #             cell = ws2.cell(1, col, text)
# #             cell.font = header_font
# #             cell.fill = header_fill
# #             cell.alignment = center
# #             cell.border = border
# #         ws2.row_dimensions[1].height = 22

# #         # Data rows
# #         for i, sample in enumerate(samples, start=2):
# #             ws2.cell(i, 1, i - 1).alignment = center
# #             ws2.cell(i, 1).border = border
# #             ws2.cell(i, 2, sample).alignment = center
# #             ws2.cell(i, 2).border = border
# #             ws2.row_dimensions[i].height = 18

# #         # Line chart
# #         chart = LineChart()
# #         chart.title = "Throughput Over Time"
# #         chart.style = 10
# #         chart.y_axis.title = "msg/s"
# #         chart.x_axis.title = "Second"
# #         chart.width = 22
# #         chart.height = 14

# #         data = Reference(ws2, min_col=2, min_row=1, max_row=len(samples) + 1)
# #         chart.add_data(data, titles_from_data=True)
# #         cats = Reference(ws2, min_col=1, min_row=2, max_row=len(samples) + 1)
# #         chart.set_categories(cats)
# #         ws2.add_chart(chart, "D2")

# #     # ── Stream to browser ────────────────────────────────────────
# #     output = io.BytesIO()
# #     wb.save(output)
# #     output.seek(0)

# #     filename = f"solengineer-perf-{results.get('target','test').replace('/','_')}.xlsx"

# #     return StreamingResponse(
# #         output,
# #         media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
# #         headers={"Content-Disposition": f"attachment; filename={filename}"}
# #     )


# # # ─── WebSocket Perf Test ─────────────────────────────────────────────────────

# # @router.websocket("/ws/perf")
# # async def perf_ws(websocket: WebSocket):
# #     global _current_tester
# #     await websocket.accept()

# #     service = state.get_service()
# #     if not service:
# #         await websocket.send_json({"error": "Not connected to broker"})
# #         await websocket.close()
# #         return

# #     try:
# #         config = await asyncio.wait_for(websocket.receive_text(), timeout=10)
# #         req = json.loads(config)
# #     except Exception:
# #         await websocket.close()
# #         return

# #     tester = PerfTester(service)
# #     _current_tester = tester
# #     results_holder = {}
# #     loop = asyncio.get_event_loop()

# #     def on_progress(received, total, elapsed, tps):
# #         asyncio.run_coroutine_threadsafe(
# #             websocket.send_json({
# #                 "type": "progress",
# #                 "received": received,
# #                 "total": total,
# #                 "elapsed": elapsed,
# #                 "tps": tps,
# #                 "percent": round(received / total * 100, 1)
# #             }),
# #             loop
# #         )

# #     def run_test():
# #         try:
# #             results_holder["data"] = tester.run(
# #                 target=req.get("target", "perf/test"),
# #                 target_type=req.get("target_type", "topic"),
# #                 message_count=req.get("message_count", 1000),
# #                 message_size=req.get("message_size", 256),
# #                 message_template=req.get("message_template", ""),
# #                 delivery_mode=req.get("delivery_mode", "direct"),
# #                 rate_limit=req.get("rate_limit", 0),
# #                 window_size=req.get("window_size", 50),
# #                 queue_type=req.get("queue_type", "exclusive"), #new 
# #                 warmup_count=req.get("warmup_count", 0),# new 
# #                 consumer_count=req.get("consumer_count", 1), #new
# #                 on_progress=on_progress
# #             )
# #         except Exception as e:
# #             results_holder["error"] = str(e)

# #     thread = threading.Thread(target=run_test)
# #     thread.start()

# #     try:
# #         while thread.is_alive():
# #             await asyncio.sleep(0.3)
# #         thread.join()

# #         if "error" in results_holder:
# #             await websocket.send_json({
# #                 "type": "error",
# #                 "error": results_holder["error"]
# #             })
# #         elif "data" in results_holder:
# #             await websocket.send_json({
# #                 "type": "complete",
# #                 "results": results_holder["data"],
# #                 "timestamp": datetime.now().isoformat()
# #             })
# #     except WebSocketDisconnect:
# #         tester.stop()
# #     finally:
# #         try:
# #             await websocket.close()
# #         except Exception:
# #             pass


from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from solengineer.api import state
from solengineer.perf.perf_tester import PerfTester
import asyncio
import threading
import json
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from datetime import datetime

router = APIRouter()

_current_tester: PerfTester | None = None


class PerfRequest(BaseModel):
    target: str
    target_type: str = "topic"
    publish_topic: str = ""         # ← NEW
    message_count: int = 1000
    message_size: int = 256
    message_template: str = ""
    delivery_mode: str = "direct"
    rate_limit: int = 0
    window_size: int = 50
    queue_type: str = "exclusive"
    warmup_count: int = 0
    consumer_count: int = 1


@router.post("/api/perf")
def run_perf(req: PerfRequest):
    service = state.get_service()
    if not service:
        raise HTTPException(status_code=400, detail="Not connected to broker")
    tester = PerfTester(service)
    results = tester.run(
        target=req.target,
        target_type=req.target_type,
        publish_topic=req.publish_topic,
        message_count=req.message_count,
        message_size=req.message_size,
        message_template=req.message_template,
        delivery_mode=req.delivery_mode,
        rate_limit=req.rate_limit,
        window_size=req.window_size,
        queue_type=req.queue_type,
        warmup_count=req.warmup_count,
        consumer_count=req.consumer_count
    )
    return results


@router.post("/api/perf/stop")
def stop_perf():
    global _current_tester
    if _current_tester:
        _current_tester.stop()
    return {"status": "stopped"}


# ─── Excel Export ─────────────────────────────────────────────────────────────

@router.post("/api/perf/export/excel")
def export_perf_excel(results: dict):
    wb = openpyxl.Workbook()

    header_font  = Font(bold=True, color="FFFFFF", size=11)
    header_fill  = PatternFill("solid", fgColor="1A73E8")
    label_font   = Font(bold=True, color="5A6675", size=10)
    value_font   = Font(color="1A2733", size=10)
    title_font   = Font(bold=True, color="1A2733", size=13)
    center       = Alignment(horizontal="center", vertical="center")
    left         = Alignment(horizontal="left", vertical="center")
    thin         = Side(style="thin", color="D4DAE0")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 32

    ws.merge_cells("A1:B1")
    ws["A1"] = "SolEngineer — Performance Test Results"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws["A1"].fill = PatternFill("solid", fgColor="F5F7FA")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:B2")
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font = Font(color="8A96A3", size=9, italic=True)
    ws["A2"].alignment = center
    ws.row_dimensions[2].height = 18

    def write_section(start_row, title, rows):
        ws.merge_cells(f"A{start_row}:B{start_row}")
        ws[f"A{start_row}"] = title
        ws[f"A{start_row}"].font = header_font
        ws[f"A{start_row}"].fill = header_fill
        ws[f"A{start_row}"].alignment = center
        ws[f"B{start_row}"].fill = header_fill
        ws.row_dimensions[start_row].height = 22
        r = start_row + 1
        for label, value in rows:
            ws.cell(r, 1, label).font = label_font
            ws.cell(r, 1).alignment = left
            ws.cell(r, 1).border = border
            ws.cell(r, 2, value).font = value_font
            ws.cell(r, 2).alignment = left
            ws.cell(r, 2).border = border
            ws.row_dimensions[r].height = 20
            r += 1
        return r

    next_row = write_section(4, "Configuration", [
        ("Target",               results.get("target", "")),
        ("Target Type",          results.get("target_type", "").capitalize()),
        ("Queue Type",           results.get("queue_type") or "N/A"),
        ("Consumer Count",       results.get("consumer_count", 1)),
        ("Delivery Mode",        results.get("delivery_mode", "").capitalize()),
        ("Message Count",        results.get("message_count", "")),
        ("Message Size (bytes)", results.get("message_size_bytes", "")),
        ("Rate Limit (msg/s)",   results.get("rate_limit", 0) or "Unlimited"),
        ("Window Size",          results.get("window_size") or "N/A (Direct)"),
        ("Warmup Messages",      results.get("warmup_count", 0)),
    ])

    next_row += 1
    write_section(next_row, "Results", [
        ("Throughput (msg/s)",  results.get("throughput_msg_per_sec", "")),
        ("Avg Latency (ms)",    results.get("avg_latency_ms", "")),
        ("P95 Latency (ms)",    results.get("p95_latency_ms", "")),
        ("P99 Latency (ms)",    results.get("p99_latency_ms", "")),
        ("Min Latency (ms)",    results.get("min_latency_ms", "")),
        ("Max Latency (ms)",    results.get("max_latency_ms", "")),
        ("Warmup Avg (ms)",     results.get("warmup_avg_ms") or "N/A"),
        ("Warmup Max (ms)",     results.get("warmup_max_ms") or "N/A"),
        ("Messages Sent",       results.get("sent", "")),
        ("Messages Received",   results.get("received", "")),
        ("Errors",              results.get("errors", "")),
        ("Total Time (sec)",    results.get("total_time_sec", "")),
    ])

    # Sheet 2 — Egress (receive rate)
    egress = results.get("throughput_samples", [])
    if egress:
        ws2 = wb.create_sheet("Egress (Receive Rate)")
        ws2.sheet_view.showGridLines = False
        ws2.column_dimensions["A"].width = 15
        ws2.column_dimensions["B"].width = 25
        for col, text in [(1, "Second"), (2, "Egress (msg/s)")]:
            cell = ws2.cell(1, col, text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border
        ws2.row_dimensions[1].height = 22
        for i, v in enumerate(egress, start=2):
            ws2.cell(i, 1, i - 1).alignment = center
            ws2.cell(i, 1).border = border
            ws2.cell(i, 2, v).alignment = center
            ws2.cell(i, 2).border = border
        chart2 = LineChart()
        chart2.title = "Egress — Receive Rate"
        chart2.style = 10
        chart2.y_axis.title = "msg/s"
        chart2.x_axis.title = "Second"
        chart2.width = 22
        chart2.height = 14
        data2 = Reference(ws2, min_col=2, min_row=1, max_row=len(egress) + 1)
        chart2.add_data(data2, titles_from_data=True)
        ws2.add_chart(chart2, "D2")

    # Sheet 3 — Ingress (publish rate)
    ingress = results.get("ingress_samples", [])
    if ingress:
        ws3 = wb.create_sheet("Ingress (Publish Rate)")
        ws3.sheet_view.showGridLines = False
        ws3.column_dimensions["A"].width = 15
        ws3.column_dimensions["B"].width = 25
        for col, text in [(1, "Second"), (2, "Ingress (msg/s)")]:
            cell = ws3.cell(1, col, text)
            cell.font = header_font
            cell.fill = PatternFill("solid", fgColor="00874A")
            cell.alignment = center
            cell.border = border
        ws3.row_dimensions[1].height = 22
        for i, v in enumerate(ingress, start=2):
            ws3.cell(i, 1, i - 1).alignment = center
            ws3.cell(i, 1).border = border
            ws3.cell(i, 2, v).alignment = center
            ws3.cell(i, 2).border = border
        chart3 = LineChart()
        chart3.title = "Ingress — Publish Rate"
        chart3.style = 10
        chart3.y_axis.title = "msg/s"
        chart3.x_axis.title = "Second"
        chart3.width = 22
        chart3.height = 14
        data3 = Reference(ws3, min_col=2, min_row=1, max_row=len(ingress) + 1)
        chart3.add_data(data3, titles_from_data=True)
        ws3.add_chart(chart3, "D2")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"solengineer-perf-{results.get('target','test').replace('/','_')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ─── WebSocket Perf ───────────────────────────────────────────────────────────

@router.websocket("/ws/perf")
async def perf_ws(websocket: WebSocket):
    global _current_tester
    await websocket.accept()

    service = state.get_service()
    if not service:
        await websocket.send_json({"error": "Not connected to broker"})
        await websocket.close()
        return

    try:
        config = await asyncio.wait_for(websocket.receive_text(), timeout=10)
        req = json.loads(config)
    except Exception:
        await websocket.close()
        return

    tester = PerfTester(service)
    _current_tester = tester
    results_holder = {}
    loop = asyncio.get_event_loop()

    def on_progress(received, total, elapsed, tps, ingress_tps=0):
        asyncio.run_coroutine_threadsafe(
            websocket.send_json({
                "type":        "progress",
                "received":    received,
                "total":       total,
                "elapsed":     elapsed,
                "tps":         tps,
                "ingress_tps": ingress_tps,
                "percent":     round(received / total * 100, 1)
            }),
            loop
        )

    def run_test():
        try:
            results_holder["data"] = tester.run(
                target=req.get("target", "perf/test"),
                target_type=req.get("target_type", "topic"),
                publish_topic=req.get("publish_topic", ""),
                message_count=req.get("message_count", 1000),
                message_size=req.get("message_size", 256),
                message_template=req.get("message_template", ""),
                delivery_mode=req.get("delivery_mode", "direct"),
                rate_limit=req.get("rate_limit", 0),
                window_size=req.get("window_size", 50),
                queue_type=req.get("queue_type", "exclusive"),
                warmup_count=req.get("warmup_count", 0),
                consumer_count=req.get("consumer_count", 1),
                on_progress=on_progress
            )
        except Exception as e:
            results_holder["error"] = str(e)

    thread = threading.Thread(target=run_test)
    thread.start()

    try:
        while thread.is_alive():
            await asyncio.sleep(0.3)
        thread.join()

        if "error" in results_holder:
            await websocket.send_json({
                "type": "error",
                "error": results_holder["error"]
            })
        elif "data" in results_holder:
            await websocket.send_json({
                "type":      "complete",
                "results":   results_holder["data"],
                "timestamp": datetime.now().isoformat()
            })
    except WebSocketDisconnect:
        tester.stop()
    finally:
        try:
            await websocket.close()
        except Exception:
            pass